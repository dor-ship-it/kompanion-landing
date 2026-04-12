"""
KompanionAI Document Generator – Flask Backend
Run: python app.py
Requires: ANTHROPIC_API_KEY in environment or .env file
"""

import os
import io
import json
import re
import datetime

from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
import anthropic

app = Flask(__name__, static_folder='.', static_url_path='')
CORS(app)

# Load .env if present
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

client = anthropic.Anthropic(api_key=os.environ.get('ANTHROPIC_API_KEY', ''))

# ──────────────────────────────────────────────
# Brand Colors (KompanionAI)
# ──────────────────────────────────────────────
YELLOW   = '#FFD039'
DARK     = '#0E0D0B'
DARK2    = '#181714'
CREAM    = '#FFFBF0'
SALMON   = '#FF8C69'


# ──────────────────────────────────────────────
# Routes
# ──────────────────────────────────────────────

@app.route('/')
def index():
    return app.send_static_file('generator.html')


@app.route('/api/generate', methods=['POST'])
def api_generate():
    data = request.get_json(force=True)
    fmt  = data.get('format', 'pdf')          # pdf | slides | word | pptx | image
    desc = (data.get('description') or '').strip()

    if not desc:
        return jsonify({'error': 'Description is required'}), 400

    if not os.environ.get('ANTHROPIC_API_KEY'):
        return jsonify({'error': 'ANTHROPIC_API_KEY is not set. Add it to .env file.'}), 500

    try:
        content = generate_content(desc, fmt)
        if fmt == 'pdf':
            return make_pdf_report(content)
        elif fmt == 'slides':
            return make_pdf_slides(content)
        elif fmt == 'word':
            return make_word(content)
        elif fmt == 'pptx':
            return make_pptx(content)
        elif fmt == 'image':
            return make_image(content)
        else:
            return jsonify({'error': f'Unknown format: {fmt}'}), 400

    except json.JSONDecodeError as e:
        return jsonify({'error': f'AI returned invalid JSON: {e}'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ──────────────────────────────────────────────
# AI Content Generation
# ──────────────────────────────────────────────

def generate_content(description: str, fmt: str) -> dict:
    today = datetime.date.today().strftime('%d.%m.%Y')
    is_slides = fmt in ('slides', 'pptx')

    if is_slides:
        prompt = f"""Create a professional business presentation based on this request: "{description}"

Return ONLY a valid JSON object – no explanation, no markdown, no code fences.

JSON schema:
{{
  "title": "string",
  "subtitle": "string",
  "date": "{today}",
  "slides": [
    {{"type": "title",   "heading": "string", "subheading": "string"}},
    {{"type": "content", "heading": "string", "bullets": ["string", "string", "string"]}},
    {{"type": "stat",    "heading": "string", "stats": [{{"number": "string", "label": "string"}}]}},
    {{"type": "summary", "heading": "string", "content": "string"}}
  ]
}}

Rules:
- Create 6-8 slides total (title + content + optional stat + summary)
- Use the same language as the request
- Be specific, data-driven, and professional
- Return ONLY raw JSON"""

    else:
        prompt = f"""Create a professional business report based on this request: "{description}"

Return ONLY a valid JSON object – no explanation, no markdown, no code fences.

JSON schema:
{{
  "title": "string",
  "subtitle": "string",
  "date": "{today}",
  "sections": [
    {{
      "heading": "string",
      "content": "string",
      "bullets": ["string", "string", "string"]
    }}
  ],
  "summary": "string"
}}

Rules:
- Create 4-6 sections
- Use the same language as the request
- Be specific, professional, and insightful
- Return ONLY raw JSON"""

    message = client.messages.create(
        model='claude-opus-4-6',
        max_tokens=4096,
        messages=[{'role': 'user', 'content': prompt}]
    )

    raw = message.content[0].text.strip()

    # Strip markdown fences if present
    raw = re.sub(r'^```(?:json)?\s*', '', raw)
    raw = re.sub(r'\s*```$', '', raw)

    return json.loads(raw)


# ──────────────────────────────────────────────
# PDF Report
# ──────────────────────────────────────────────

def make_pdf_report(content: dict):
    from weasyprint import HTML as WP

    sections_html = ''
    for s in content.get('sections', []):
        bullets = ''.join(f'<li>{b}</li>' for b in s.get('bullets', []))
        sections_html += f"""
        <div class="section">
          <h2>{s.get('heading','')}</h2>
          <p>{s.get('content','')}</p>
          {'<ul>' + bullets + '</ul>' if bullets else ''}
        </div>"""

    html = f"""<!DOCTYPE html>
<html>
<head><meta charset="UTF-8">
<style>
@page {{ size: A4; margin: 0; }}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: Arial, Helvetica, sans-serif; direction: rtl; background: white; color: #1a1a1a; }}

.top-bar {{ background: {YELLOW}; height: 6px; }}
.header {{ background: {DARK}; padding: 24px 48px; display: flex; align-items: center; gap: 14px; }}
.logo-k {{ background: {YELLOW}; width: 36px; height: 36px; border-radius: 8px; display: inline-flex; align-items: center; justify-content: center; font-weight: 900; color: {DARK}; font-size: 18px; }}
.brand {{ color: {YELLOW}; font-size: 20px; font-weight: 800; }}
.hero {{ background: {DARK}; color: white; padding: 44px 48px 36px; border-bottom: 3px solid {YELLOW}; }}
.hero h1 {{ font-size: 34px; font-weight: 900; color: {YELLOW}; margin-bottom: 10px; }}
.hero .sub {{ font-size: 16px; color: #aaa; }}
.hero .date {{ font-size: 13px; color: #555; margin-top: 14px; }}
.content {{ padding: 40px 48px; }}
.section {{ margin-bottom: 36px; page-break-inside: avoid; }}
.section h2 {{ font-size: 19px; font-weight: 800; color: {DARK}; border-right: 4px solid {YELLOW}; padding-right: 14px; margin-bottom: 12px; }}
.section p {{ font-size: 14px; color: #444; line-height: 1.8; margin-bottom: 12px; }}
.section ul {{ padding-right: 20px; }}
.section li {{ font-size: 14px; color: #555; margin: 6px 0; line-height: 1.6; }}
.summary {{ background: {DARK}; color: white; padding: 36px 48px; }}
.summary h2 {{ color: {YELLOW}; font-size: 18px; font-weight: 800; margin-bottom: 12px; }}
.summary p {{ font-size: 14px; color: #ccc; line-height: 1.8; }}
.footer {{ background: #111; color: #444; text-align: center; padding: 16px; font-size: 11px; }}
</style>
</head>
<body>
<div class="top-bar"></div>
<div class="header">
  <div class="logo-k">K</div>
  <span class="brand">KompanionAI</span>
</div>
<div class="hero">
  <h1>{content.get('title','דוח')}</h1>
  <div class="sub">{content.get('subtitle','')}</div>
  <div class="date">{content.get('date','')}</div>
</div>
<div class="content">{sections_html}</div>
<div class="summary">
  <h2>סיכום מנהלים</h2>
  <p>{content.get('summary','')}</p>
</div>
<div class="footer">KompanionAI · Generated with AI · {content.get('date','')}</div>
</body>
</html>"""

    pdf = WP(string=html).write_pdf()
    ts  = datetime.date.today().strftime('%Y%m%d')
    return send_file(io.BytesIO(pdf), mimetype='application/pdf',
                     as_attachment=True, download_name=f'kompanion-report-{ts}.pdf')


# ──────────────────────────────────────────────
# PDF Slides
# ──────────────────────────────────────────────

def make_pdf_slides(content: dict):
    from weasyprint import HTML as WP

    slides_html = ''
    for i, slide in enumerate(content.get('slides', [])):
        pb = '' if i == 0 else 'page-break-before:always;'
        t  = slide.get('type', 'content')

        if t == 'title':
            slides_html += f"""
<div class="slide" style="{pb}">
  <div class="slide-center">
    <div class="slide-logo-row">
      <span class="slide-logo-k">K</span>
      <span class="slide-brand">KompanionAI</span>
    </div>
    <h1 class="title-heading">{slide.get('heading','')}</h1>
    <p class="title-sub">{slide.get('subheading','')}</p>
    <span class="slide-date">{content.get('date','')}</span>
  </div>
  <div class="bottom-bar"></div>
</div>"""

        elif t == 'stat':
            stats_html = ''.join(
                f'<div class="stat-box"><div class="stat-num">{s.get("number","")}</div>'
                f'<div class="stat-lbl">{s.get("label","")}</div></div>'
                for s in slide.get('stats', [])
            )
            slides_html += f"""
<div class="slide" style="{pb}">
  <div class="slide-inner">
    <h2 class="slide-h2">{slide.get('heading','')}</h2>
    <div class="divider"></div>
    <div class="stats-row">{stats_html}</div>
  </div>
  <div class="slide-num">{i+1}</div>
  <div class="bottom-bar"></div>
</div>"""

        elif t == 'summary':
            slides_html += f"""
<div class="slide" style="{pb}">
  <div class="slide-inner">
    <h2 class="slide-h2">{slide.get('heading','')}</h2>
    <div class="divider"></div>
    <p class="summary-p">{slide.get('content','')}</p>
  </div>
  <div class="slide-num">{i+1}</div>
  <div class="bottom-bar"></div>
</div>"""

        else:
            bullets = ''.join(f'<li>{b}</li>' for b in slide.get('bullets', []))
            slides_html += f"""
<div class="slide" style="{pb}">
  <div class="slide-inner">
    <h2 class="slide-h2">{slide.get('heading','')}</h2>
    <div class="divider"></div>
    <ul class="slide-bullets">{bullets}</ul>
  </div>
  <div class="slide-num">{i+1}</div>
  <div class="bottom-bar"></div>
</div>"""

    html = f"""<!DOCTYPE html>
<html>
<head><meta charset="UTF-8">
<style>
@page {{ size: 297mm 167mm; margin: 0; }}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: Arial, Helvetica, sans-serif; direction: rtl; }}

.slide {{ width: 297mm; height: 167mm; background: {DARK}; color: white; position: relative; overflow: hidden; display: flex; align-items: center; justify-content: center; }}
.bottom-bar {{ position: absolute; bottom: 0; left: 0; right: 0; height: 5px; background: {YELLOW}; }}
.slide-num {{ position: absolute; bottom: 14px; left: 28px; font-size: 11px; color: #444; }}

.slide-center {{ text-align: center; padding: 40px; }}
.slide-logo-row {{ display: flex; align-items: center; justify-content: center; gap: 10px; margin-bottom: 28px; }}
.slide-logo-k {{ background: {YELLOW}; width: 36px; height: 36px; border-radius: 8px; display: inline-flex; align-items: center; justify-content: center; font-weight: 900; color: {DARK}; font-size: 18px; }}
.slide-brand {{ color: {YELLOW}; font-size: 18px; font-weight: 800; }}
.title-heading {{ font-size: 46px; font-weight: 900; color: white; margin-bottom: 16px; line-height: 1.15; }}
.title-sub {{ font-size: 20px; color: #aaa; margin-bottom: 20px; }}
.slide-date {{ font-size: 13px; color: #555; }}

.slide-inner {{ padding: 36px 56px; width: 100%; }}
.slide-h2 {{ font-size: 28px; font-weight: 900; color: {YELLOW}; margin-bottom: 14px; }}
.divider {{ height: 2px; background: rgba(255,208,57,0.25); margin-bottom: 20px; }}

.slide-bullets {{ padding-right: 24px; list-style: none; }}
.slide-bullets li {{ font-size: 17px; color: #ddd; margin: 10px 0; line-height: 1.5; padding-right: 16px; position: relative; }}
.slide-bullets li::before {{ content: "▸"; color: {YELLOW}; position: absolute; right: 0; }}

.stats-row {{ display: flex; gap: 36px; justify-content: center; margin-top: 16px; flex-wrap: wrap; }}
.stat-box {{ text-align: center; }}
.stat-num {{ font-size: 56px; font-weight: 900; color: {YELLOW}; }}
.stat-lbl {{ font-size: 15px; color: #aaa; margin-top: 6px; }}

.summary-p {{ font-size: 18px; color: #ccc; line-height: 1.8; }}
</style>
</head>
<body>{slides_html}</body>
</html>"""

    pdf = WP(string=html).write_pdf()
    ts  = datetime.date.today().strftime('%Y%m%d')
    return send_file(io.BytesIO(pdf), mimetype='application/pdf',
                     as_attachment=True, download_name=f'kompanion-slides-{ts}.pdf')


# ──────────────────────────────────────────────
# Word DOCX
# ──────────────────────────────────────────────

def make_word(content: dict):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)

    Y  = RGBColor(0xFF, 0xD0, 0x39)
    DK = RGBColor(0x0E, 0x0D, 0x0B)
    GR = RGBColor(0x44, 0x44, 0x44)
    LG = RGBColor(0x88, 0x88, 0x88)

    def add_para(text, size, bold=False, color=DK, align=WD_ALIGN_PARAGRAPH.RIGHT):
        p = doc.add_paragraph()
        p.alignment = align
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        return p

    add_para(content.get('title', ''), 28, bold=True, color=Y, align=WD_ALIGN_PARAGRAPH.CENTER)
    if content.get('subtitle'):
        add_para(content['subtitle'], 16, color=GR, align=WD_ALIGN_PARAGRAPH.CENTER)
    if content.get('date'):
        add_para(content['date'], 11, color=LG, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()

    for s in content.get('sections', []):
        add_para(s.get('heading', ''), 16, bold=True, color=DK)
        if s.get('content'):
            add_para(s['content'], 12, color=GR)
        for b in s.get('bullets', []):
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(b)
            r.font.size = Pt(12)
            r.font.color.rgb = GR
        doc.add_paragraph()

    if content.get('summary'):
        add_para('סיכום מנהלים', 16, bold=True, color=Y)
        add_para(content['summary'], 12, color=GR)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    ts = datetime.date.today().strftime('%Y%m%d')
    return send_file(buf,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                     as_attachment=True, download_name=f'kompanion-report-{ts}.docx')


# ──────────────────────────────────────────────
# PowerPoint PPTX
# ──────────────────────────────────────────────

def make_pptx(content: dict):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    W, H = Inches(13.33), Inches(7.5)
    Y  = RGBColor(0xFF, 0xD0, 0x39)
    DK = RGBColor(0x0E, 0x0D, 0x0B)
    WH = RGBColor(0xFF, 0xFF, 0xFF)
    GR = RGBColor(0xAA, 0xAA, 0xAA)

    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    def bg(slide, color=DK):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def rect(slide, l, t, w, h, color):
        from pptx.util import Inches
        shp = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        return shp

    def txt(slide, text, l, t, w, h, size, bold=False, color=WH, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p  = tf.paragraphs[0]
        p.alignment = align
        rn = p.add_run()
        rn.text = text
        rn.font.size = Pt(size)
        rn.font.bold = bold
        rn.font.color.rgb = color

    def yellow_bar(slide):
        rect(slide, 0, 7.42, 13.33, 0.08, Y)

    for slide_data in content.get('slides', []):
        sl = prs.slides.add_slide(blank)
        bg(sl)
        yellow_bar(sl)
        t = slide_data.get('type', 'content')

        if t == 'title':
            rect(sl, 5.7, 1.4, 0.45, 0.45, Y)
            txt(sl, 'KompanionAI', 6.25, 1.45, 4, 0.45, 18, bold=True, color=Y)
            txt(sl, slide_data.get('heading', ''), 1, 2.5, 11.3, 1.6, 44, bold=True, color=WH, align=PP_ALIGN.CENTER)
            txt(sl, slide_data.get('subheading', ''), 1.5, 4.3, 10.3, 0.8, 22, color=GR, align=PP_ALIGN.CENTER)
            txt(sl, content.get('date', ''), 5.5, 5.8, 2.5, 0.4, 12, color=RGBColor(0x55,0x55,0x55), align=PP_ALIGN.CENTER)

        elif t == 'stat':
            txt(sl, slide_data.get('heading', ''), 0.5, 0.5, 12, 0.8, 28, bold=True, color=Y)
            rect(sl, 0.5, 1.45, 12.33, 0.03, RGBColor(0x33,0x30,0x20))
            stats = slide_data.get('stats', [])
            cw = 12.0 / max(len(stats), 1)
            for i, s in enumerate(stats):
                x = 0.5 + i * cw
                txt(sl, s.get('number', ''), x, 2.2, cw, 1.4, 54, bold=True, color=Y, align=PP_ALIGN.CENTER)
                txt(sl, s.get('label', ''),  x, 3.8, cw, 0.5, 16, color=GR, align=PP_ALIGN.CENTER)

        else:
            txt(sl, slide_data.get('heading', ''), 0.5, 0.5, 12, 0.8, 28, bold=True, color=Y)
            rect(sl, 0.5, 1.45, 12.33, 0.03, RGBColor(0x33,0x30,0x20))
            if slide_data.get('bullets'):
                y = 1.7
                for b in slide_data['bullets']:
                    txt(sl, f'▸  {b}', 0.6, y, 12, 0.65, 18, color=WH)
                    y += 0.78
            elif slide_data.get('content'):
                txt(sl, slide_data['content'], 0.6, 1.7, 12, 4.5, 18, color=RGBColor(0xCC,0xCC,0xCC))

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    ts = datetime.date.today().strftime('%Y%m%d')
    return send_file(buf,
                     mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                     as_attachment=True, download_name=f'kompanion-slides-{ts}.pptx')


# ──────────────────────────────────────────────
# PNG Image (infographic)
# ──────────────────────────────────────────────

def make_image(content: dict):
    from weasyprint import HTML as WP

    title = content.get('title', 'KompanionAI')
    items = content.get('sections') or content.get('slides', [])

    cards_html = ''
    for item in items[:4]:
        heading = item.get('heading', '')
        text = item.get('content', '') or ', '.join(item.get('bullets', [])[:3])
        text = (text[:160] + '…') if len(text) > 160 else text
        cards_html += f"""
<div class="card">
  <div class="card-title">{heading}</div>
  <div class="card-text">{text}</div>
</div>"""

    html = f"""<!DOCTYPE html>
<html>
<head><meta charset="UTF-8">
<style>
@page {{ size: 1200px 675px; margin: 0; }}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: Arial, Helvetica, sans-serif; width: 1200px; height: 675px; background: {DARK}; color: white; overflow: hidden; direction: rtl; }}
.top-bar {{ background: {YELLOW}; height: 6px; }}
.header {{ padding: 22px 40px 14px; display: flex; align-items: center; justify-content: space-between; }}
.logo-row {{ display: flex; align-items: center; gap: 12px; }}
.logo-k {{ background: {YELLOW}; width: 34px; height: 34px; border-radius: 8px; display: inline-flex; align-items: center; justify-content: center; font-weight: 900; color: {DARK}; font-size: 17px; }}
.brand {{ color: {YELLOW}; font-size: 18px; font-weight: 800; }}
.date-tag {{ font-size: 12px; color: #444; }}
.main-title {{ font-size: 30px; font-weight: 900; color: white; padding: 4px 40px 20px; line-height: 1.2; }}
.main-title span {{ color: {YELLOW}; }}
.cards {{ display: grid; grid-template-columns: 1fr 1fr; gap: 14px; padding: 0 40px; }}
.card {{ background: rgba(255,255,255,0.05); border: 1px solid rgba(255,208,57,0.18); border-radius: 14px; padding: 20px 22px; }}
.card-title {{ font-size: 14px; font-weight: 800; color: {YELLOW}; margin-bottom: 8px; }}
.card-text {{ font-size: 13px; color: #aaa; line-height: 1.6; }}
.bottom-bar {{ position: absolute; bottom: 0; left: 0; right: 0; background: {YELLOW}; height: 5px; }}
</style>
</head>
<body>
<div class="top-bar"></div>
<div class="header">
  <div class="logo-row">
    <div class="logo-k">K</div>
    <span class="brand">KompanionAI</span>
  </div>
  <div class="date-tag">{content.get('date','')}</div>
</div>
<div class="main-title">{title}<br><span>{content.get('subtitle','')}</span></div>
<div class="cards">{cards_html}</div>
<div class="bottom-bar"></div>
</body>
</html>"""

    pdf_bytes = WP(string=html).write_pdf()
    ts = datetime.date.today().strftime('%Y%m%d')

    # Try converting PDF page → PNG
    try:
        from pdf2image import convert_from_bytes
        imgs = convert_from_bytes(pdf_bytes, dpi=144, first_page=1, last_page=1)
        buf = io.BytesIO()
        imgs[0].save(buf, format='PNG')
        buf.seek(0)
        return send_file(buf, mimetype='image/png',
                         as_attachment=True, download_name=f'kompanion-image-{ts}.png')
    except Exception:
        # Fallback: return as PDF (user can screenshot/export)
        return send_file(io.BytesIO(pdf_bytes), mimetype='application/pdf',
                         as_attachment=True,
                         download_name=f'kompanion-image-{ts}.pdf')


# ──────────────────────────────────────────────

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5001))
    print(f'\n🟡  KompanionAI Document Generator')
    print(f'    http://localhost:{port}/generator.html\n')
    app.run(debug=True, port=port)
